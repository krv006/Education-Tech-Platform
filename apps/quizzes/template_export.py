"""Test tuzish uchun bo'sh shablon fayl generatsiya qilish (.docx / .xlsx).

`docx_import.py` va `xlsx_import.py` parserlari kutgan format bilan AYNAN
mos — shuning uchun shablonni to'ldirib qaytadan import qilganda muammo
bo'lmaydi (round-trip: har ikkalasi ham `apps.quizzes.tests`da tekshiriladi).

Savollar soni frontenddagi "Nechta savol?" maydoni bilan bir xil — o'qituvchi
nechta savol so'rasa, shablonda shuncha bo'sh savol bloki/qatori bo'ladi
(xuddi ilova ichidagi bo'sh shablon generatori kabi, faqat fayl holida).
"""
from io import BytesIO

MAX_TEMPLATE_QUESTIONS = 100
DEFAULT_OPTION_COUNT = 4
_OPTION_LETTERS = 'ABCDE'

_QUESTION_PLACEHOLDER = 'Savol matnini shu yerga yozing'
_ANSWER_PLACEHOLDER = "A/B/C/D dan birini yozing"


def clamp_count(raw) -> int:
    try:
        count = int(raw)
    except (TypeError, ValueError):
        count = 10
    return max(1, min(MAX_TEMPLATE_QUESTIONS, count))


def build_docx_template(count: int, option_count: int = DEFAULT_OPTION_COUNT) -> bytes:
    import docx

    document = docx.Document()
    for i in range(1, count + 1):
        document.add_paragraph(f'{i}. [{_QUESTION_PLACEHOLDER}]')
        for letter in _OPTION_LETTERS[:option_count]:
            document.add_paragraph(f'{letter}) [{letter}-variant matni]')
        document.add_paragraph(f"To'g'ri javob: [{_ANSWER_PLACEHOLDER}]")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx_template(count: int) -> bytes:
    """Ustunlar soni doim FIKS (Savol + A..E + To'g'ri javob = 7 ustun) —
    `xlsx_import.py`dagi `_OPTION_COLUMNS`/`_ANSWER_COLUMN` ustun-pozitsiyaga
    bog'liq bo'lgani uchun (Excel jadvali .docx'dan farqli, moslashuvchan
    ustun soniga ega bo'lolmaydi). Kerak bo'lmagan variant ustuni (masalan E)
    shunchaki bo'sh qoldirilishi mumkin — bu xato emas."""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Test savollari'
    sheet.append(['Savol', *_OPTION_LETTERS, "To'g'ri javob"])
    for _ in range(count):
        sheet.append([
            f'[{_QUESTION_PLACEHOLDER}]',
            *[f'[{letter}-variant matni]' for letter in _OPTION_LETTERS],
            f'[{_ANSWER_PLACEHOLDER}]',
        ])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
